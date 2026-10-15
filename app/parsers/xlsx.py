"""Парсер XLSX: openpyxl, строки спецификации → чанки «артикул: ключ значение»."""
from openpyxl import load_workbook

from app.parsers.pdf import Block


def parse_xlsx(data: bytes, sheet: str | None = None) -> list[Block]:
    import io

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows = [[("" if c is None else str(c)).strip() for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()

    # Заголовок — первая непустая строка
    header_idx = next((i for i, r in enumerate(rows) if any(r)), None)
    if header_idx is None:
        return []
    header = rows[header_idx]

    blocks: list[Block] = []
    for row in rows[header_idx + 1 :]:
        if not any(row):
            continue
        pairs = {h: v for h, v in zip(header, row) if h and v}
        articul = pairs.get("Артикул", "")
        # Строка спецификации → линейный текст «артикул: ключ: значение; …»
        body = "; ".join(f"{k}: {v}" for k, v in pairs.items())
        blocks.append(
            Block(
                kind="table",
                content=f"{articul}: {body}" if articul else body,
                page=1,
                section="Спецификация",
                rows=[list(pairs.values())],
            )
        )
    return blocks
