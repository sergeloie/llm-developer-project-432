"""Парсер DOCX: python-docx, параграфы + таблицы, стили-заголовки как секции."""
from docx import Document as DocxDocument

from app.parsers.pdf import Block


def parse_docx(data: bytes) -> list[Block]:
    import io

    doc = DocxDocument(io.BytesIO(data))
    blocks: list[Block] = []
    section = ""
    page = 1  # DOCX без явной пагинации — все блоки на «странице» 1

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if "heading" in style or "заголовок" in style:
            section = text
        blocks.append(Block("text", text, page, section=section))

    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        from app.parsers.pdf import _table_to_text

        blocks.append(Block("table", _table_to_text(rows), page, section=section, rows=rows))

    return blocks
